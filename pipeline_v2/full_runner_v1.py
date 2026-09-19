"""First formal full cleaning run for the local emergency-material corpus.

The runner is intentionally an orchestration layer.  Parent/segment/title
decisions are delegated to the frozen V2.4.2.2 modules that passed the pilot
and preflight tests.  Raw files are opened read-only and are never moved,
rewritten or deleted.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata
import json
import os
import plistlib
import random
import re
import signal
import subprocess
import tempfile
import time
from collections import Counter, defaultdict
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from datasketch import MinHashLSH
from rapidfuzz.fuzz import ratio as rapidfuzz_ratio

from .classification import content_role_of
from .classification_v24 import first_effective_title
from .classification_v241 import lock_parent_type_v241
from .classification_v242 import accident_category_v242, recover_title_v242, segment_type_v242
from .classification_v2421 import classify_segment_v2421
from .deduplication import minhash
from .knowledge_extractor import extract_knowledge
from .models import PlanSection
from .models_full_v1 import (
    FullManifestRecord,
    FullParentDocument,
    FullRagChunk,
    FullSegmentRecord,
    ValidationSummary,
)
from .normalization import looks_garbled, normalize_text
from .parsers import parse_local
from .redaction import redact
from .sds_matcher import sds_from_text
from .section_splitter_v24 import audit_heading, split_sections_v24
from .title_audit_v242 import audit_title_consistency_v242
from .title_hotfix_v2422 import (
    audit_full_document_title_v2422,
    extract_sds_identity,
    recover_full_document_title_v2422,
)


PILOT_SEED = 20260824
SUPPORTED_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".csv", ".html", ".htm",
    ".txt", ".json", ".xml", ".md", ".png", ".jpg", ".jpeg", ".tif", ".tiff",
}
OUTPUT_NAMES = {
    "cleaned": "cleaned_full_v1",
    "reports": "reports_full_v1",
    "logs": "logs_full_v1",
    "manifests": "manifests_full_v1",
    "checkpoints": "checkpoints_full_v1",
}
PARENT_TYPES = {
    "government_or_regional_plan", "enterprise_plan", "enterprise_group_plan",
    "organization_plan", "sds", "regulation", "guidance_reference", "accident_case",
    "chemical_catalog", "manual_review", "reject", "other",
}
SEGMENT_TYPES = {
    "full_document", "embedded_special_section", "onsite_disposal_plan", "guidance_section",
    "accident_case_section", "reference_sentence", "list_fragment", "toc_fragment", "table",
    "appendix_sds", "embedded_sds", "appendix", "other",
}
HIGH_RISK_CONTENT_RE = re.compile(r"急救|心肺复苏|人工呼吸|医疗|消防|灭火|堵漏|中和|洗消|催吐|洗胃")
TOC_LINE_RE = re.compile(r"^.{2,120}(?:\.{4,}|…{3,}|-{5,})\s*\d{1,4}\s*$")
REG_TITLE_RE = re.compile(r"(?:规程|规定|办法|条例|标准|规范)$")


@contextmanager
def parse_deadline(seconds: int = 120):
    """Bound a single local parser call so one corrupt file cannot stall a run."""
    if not hasattr(signal, "SIGALRM"):
        yield
        return
    previous = signal.getsignal(signal.SIGALRM)

    def _timeout(_signum, _frame):
        raise TimeoutError(f"single_file_parse_timeout_after_{seconds}s")

    signal.signal(signal.SIGALRM, _timeout)
    signal.setitimer(signal.ITIMER_REAL, seconds)
    try:
        yield
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous)


def utc_now() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def stable_id(prefix: str, *parts: str, length: int = 20) -> str:
    digest = hashlib.sha256("\x1f".join(parts).encode("utf-8", errors="replace")).hexdigest()
    return f"{prefix}-{digest[:length]}"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def package_version(name: str) -> str:
    try:
        return importlib.metadata.version(name)
    except importlib.metadata.PackageNotFoundError:
        return "not-installed"


def json_dump(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def jsonl_dump(path: Path, rows: Iterable[dict[str, Any]]) -> int:
    count = 0
    with path.open("w", encoding="utf-8") as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False) + "\n")
            count += 1
    return count


def markdown_table(headers: list[str], rows: Iterable[Iterable[Any]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    for row in rows:
        clean = [str(value).replace("\n", " ").replace("|", "\\|") for value in row]
        lines.append("| " + " | ".join(clean) + " |")
    return "\n".join(lines)


def clean_filename_title(path: Path) -> str:
    value = re.sub(r"-[0-9a-f]{10,}$", "", path.stem, flags=re.I)
    value = re.sub(r"^PDF\s+", "", value, flags=re.I)
    return value.strip(" _-｜|")


def scan_raw(raw_dir: Path) -> list[dict[str, Any]]:
    """Read-only baseline scan; caller creates output directories afterwards."""
    rows: list[dict[str, Any]] = []
    for path in sorted((item for item in raw_dir.rglob("*") if item.is_file()), key=lambda item: str(item.relative_to(raw_dir))):
        stat = path.stat()
        hash_error = ""
        try:
            digest: str | None = sha256_file(path)
        except (TimeoutError, OSError) as exc:
            digest = None
            hash_error = f"sha256_unavailable:{type(exc).__name__}:{str(exc)[:200]}"
        relative = str(path.relative_to(raw_dir))
        rows.append({
            "document_id": stable_id("file", digest or "content-unavailable", relative, str(stat.st_size), str(stat.st_mtime_ns), length=24),
            "source_file": str(path),
            "relative_path": relative,
            "file_name": path.name,
            "file_extension": path.suffix.lower().lstrip("."),
            "file_size": stat.st_size,
            "sha256": digest,
            "sha256_status": "available" if digest else "content_unavailable",
            "sha256_error": hash_error,
            "modified_time": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(timespec="seconds"),
        })
    return rows


def baseline_signature(rows: list[dict[str, Any]]) -> str:
    material = "\n".join(
        f"{row['relative_path']}\t{row['file_size']}\t{row['modified_time']}\t{row.get('sha256') or 'CONTENT_UNAVAILABLE'}"
        for row in rows
    )
    return hashlib.sha256(material.encode("utf-8")).hexdigest()


def load_source_records(root: Path, raw_rows: list[dict[str, Any]]) -> tuple[dict[str, list[dict]], list[dict]]:
    source_path = root / "json" / "records.json"
    records = json.loads(source_path.read_text(encoding="utf-8")) if source_path.exists() else []
    by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in raw_rows:
        by_name[row["file_name"]].append(row)
    mapped: dict[str, list[dict]] = defaultdict(list)
    unmatched: list[dict] = []
    for index, record in enumerate(records):
        file_value = str(record.get("file") or "")
        candidates = by_name.get(Path(file_value).name, [])
        if len(candidates) == 1:
            enriched = dict(record)
            enriched["_source_record_index"] = index
            mapped[candidates[0]["document_id"]].append(enriched)
        else:
            unmatched.append({
                "source_record_index": index,
                "file": file_value,
                "title": record.get("title", ""),
                "url": record.get("url", ""),
                "reason": "no_current_raw_basename_match" if not candidates else "ambiguous_current_raw_basename",
            })
    return mapped, unmatched


def source_metadata(records: list[dict]) -> dict[str, Any]:
    urls = list(dict.fromkeys(str(row.get("url") or "") for row in records if row.get("url")))
    titles = [str(row.get("title") or "").strip() for row in records if row.get("title")]
    types = [str(row.get("type") or "").strip() for row in records if row.get("type")]
    return {
        "source_url": urls[0] if urls else "",
        "source_urls": urls,
        "source_url_status": "present" if urls else "missing",
        "source_titles": list(dict.fromkeys(titles)),
        "original_type": Counter(types).most_common(1)[0][0] if types else "",
        "source_record_count": len(records),
    }


def local_where_from_urls(path: Path) -> list[str]:
    """Read browser provenance from the existing file; never contacts the URL."""
    try:
        result = subprocess.run(
            ["/usr/bin/xattr", "-px", "com.apple.metadata:kMDItemWhereFroms", str(path)],
            check=False,
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
            timeout=3,
        )
        if result.returncode != 0:
            return []
        raw = bytes.fromhex(result.stdout.decode("ascii", errors="ignore"))
        values = plistlib.loads(raw)
        return list(dict.fromkeys(value for value in values if isinstance(value, str) and value.startswith(("http://", "https://"))))
    except (OSError, ValueError, plistlib.InvalidFileException, subprocess.TimeoutExpired):
        return []


def parse_json_text(path: Path) -> str:
    value = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    return json.dumps(value, ensure_ascii=False, indent=2)


def parse_xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"工作表：{sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                parts.append("\t".join(values))
    workbook.close()
    return "\n".join(parts)


def parse_doc_text(path: Path) -> str:
    result = subprocess.run(
        ["/usr/bin/textutil", "-convert", "txt", "-stdout", str(path)],
        check=False,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        timeout=180,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.decode("utf-8", errors="replace")[:500])
    return result.stdout.decode("utf-8", errors="replace")


def parse_csv_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "gb18030", "big5"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def parser_payload(parser: str, text: str, page_count: int = 0, warnings: list[str] | None = None) -> dict[str, Any]:
    normalized = normalize_text(text)
    chinese = len(re.findall(r"[\u4e00-\u9fff]", normalized))
    bad = len(re.findall(r"[�□]", normalized))
    return {
        "parser": parser,
        "parser_version": package_version(parser) if parser not in {"textutil", "json", "text", "openpyxl", "unsupported"} else "",
        "page_count": page_count,
        "raw_text": text,
        "normalized_text": normalized,
        "text_length": len(normalized),
        "chinese_ratio": round(chinese / max(1, len(normalized)), 4),
        "garbled_ratio": round(bad / max(1, len(normalized)), 6),
        "warnings": list(warnings or []),
        "alternatives": [],
    }


def parse_file(row: dict[str, Any], source_rows: list[dict]) -> dict[str, Any]:
    path = Path(row["source_file"])
    suffix = path.suffix.lower()
    started = time.perf_counter()
    warnings: list[str] = []
    try:
        if suffix in {".pdf", ".docx", ".html", ".htm", ".txt"}:
            parsed, alternatives = parse_local(row["document_id"], path)
            payload = parser_payload(
                parsed.parser,
                parsed.raw_text,
                len(parsed.pages),
                list(parsed.warnings),
            )
            payload["parser_version"] = parsed.parser_version
            payload["blocks"] = [item.model_dump(mode="json") for item in parsed.blocks]
            payload["alternatives"] = [
                {
                    "parser": item.parser,
                    "parser_version": item.parser_version,
                    "text_length": item.parse_quality.get("text_length", len(item.raw_text)),
                    "garbled_ratio": item.parse_quality.get("garbled_ratio", 0),
                    "warnings": item.warnings,
                }
                for item in alternatives
            ]
        elif suffix == ".doc":
            payload = parser_payload("textutil", parse_doc_text(path))
        elif suffix == ".json":
            payload = parser_payload("json", parse_json_text(path))
        elif suffix == ".xlsx":
            payload = parser_payload("openpyxl", parse_xlsx_text(path))
        elif suffix == ".xls":
            import pandas as pd

            sheets = pd.read_excel(path, sheet_name=None)
            text = "\n\n".join(f"工作表：{name}\n{frame.to_csv(index=False)}" for name, frame in sheets.items())
            payload = parser_payload("pandas", text)
        elif suffix in {".csv", ".xml", ".md"}:
            payload = parser_payload("text", parse_csv_text(path))
        elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
            payload = parser_payload("unsupported", "", warnings=["image_ocr_requires_manual_review"])
        else:
            payload = parser_payload("unsupported", "", warnings=["unsupported_file_type"])
    except Exception as exc:  # one file must never stop the full run
        payload = parser_payload("parse_failed", "", warnings=[f"parse_failed:{type(exc).__name__}:{str(exc)[:300]}"])

    if payload["text_length"] < 120 and suffix == ".pdf":
        payload["is_scanned_document"] = True
        payload["warnings"].append("scanned_pdf_or_text_layer_missing")
        payload["warnings"].append("ocr_fallback_unavailable_for_chinese_local_language_pack")
    else:
        payload["is_scanned_document"] = suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}

    if payload["text_length"] < 120 and source_rows:
        fallback = "\n\n".join(str(item.get("content") or "") for item in source_rows if item.get("content"))
        if len(fallback.strip()) >= 120:
            payload["raw_text"] = fallback
            payload["normalized_text"] = normalize_text(fallback)
            payload["text_length"] = len(payload["normalized_text"])
            payload["warnings"].append("source_record_text_fallback_used")
            payload["parser"] = "source_record_fallback"
    payload["parse_status"] = "success" if payload["text_length"] >= 120 else "failed"
    payload["processing_seconds"] = round(time.perf_counter() - started, 4)
    return payload


def choose_parent_title(path: Path, text: str, source_rows: list[dict]) -> str:
    filename = clean_filename_title(path)
    candidate = first_effective_title(text[:12000], filename)
    if candidate and len(candidate) <= 120 and not TOC_LINE_RE.search(candidate):
        return candidate
    for title in (str(row.get("title") or "").strip() for row in source_rows):
        if title and len(title) <= 100 and re.search(r"应急预案|规程|规定|办法|条例|指南|指导意见|事故通报|安全技术说明书|SDS", title, re.I):
            if not re.search(r"\.{4,}|…{3,}|现印发|负责|执行", title):
                return title.strip("《》")
    return filename


def near_duplicate_pairs(parent_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    usable = [row for row in parent_rows if len(row.get("normalized_text", "")) >= 300 and row.get("parse_status") == "success"]
    lsh = MinHashLSH(threshold=0.80, num_perm=128)
    signatures: dict[str, Any] = {}
    texts: dict[str, str] = {}
    for row in usable:
        text = row["normalized_text"]
        sample = text if len(text) <= 70000 else text[:55000] + text[-15000:]
        signature = minhash(sample)
        signatures[row["document_id"]] = signature
        texts[row["document_id"]] = sample
        lsh.insert(row["document_id"], signature)
    output: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for ident, signature in signatures.items():
        for other in lsh.query(signature):
            if other == ident:
                continue
            pair = tuple(sorted((ident, other)))
            if pair in seen:
                continue
            seen.add(pair)
            fuzzy = rapidfuzz_ratio(texts[pair[0]], texts[pair[1]]) / 100
            jac = signatures[pair[0]].jaccard(signatures[pair[1]])
            classification = "near_duplicate" if fuzzy >= 0.92 and jac >= 0.88 else "template_clone" if jac >= 0.80 else "partial_overlap"
            output.append({
                "group_id": stable_id("near", *pair),
                "canonical_record_id": pair[0],
                "duplicate_record_ids": [pair[1]],
                "document_ids": list(pair),
                "duplicate_reason": classification,
                "classification": classification,
                "metrics": {"rapidfuzz_ratio": round(fuzzy, 4), "minhash_jaccard": round(jac, 4)},
                "review_status": "pending",
            })
    return output


def exact_duplicate_groups(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: list[dict[str, Any]] = []
    by_sha: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        if row.get("sha256"):
            by_sha[row["sha256"]].append(row)
    for digest, members in by_sha.items():
        if len(members) < 2:
            continue
        ordered = sorted(members, key=lambda item: (not bool(item.get("source_url")), item["relative_path"]))
        canonical = ordered[0]["document_id"]
        groups.append({
            "duplicate_group_id": stable_id("exact", digest),
            "canonical_record_id": canonical,
            "duplicate_record_ids": [item["document_id"] for item in ordered[1:]],
            "all_source_files": [item["source_file"] for item in ordered],
            "all_source_urls": list(dict.fromkeys(url for item in ordered for url in item.get("source_urls", []) if url)),
            "duplicate_reason": "sha256_identical",
            "sha256": digest,
            "review_status": "pending",
        })
    return groups


def implicit_section(text: str, title: str, category: str) -> list[PlanSection]:
    if not text.strip():
        return []
    return [
        PlanSection(
            section_id="s000",
            heading=title or "正文",
            content=text,
            start_position=0,
            end_position=len(text),
            topic=category or "other",
            boundary_confidence=0.5,
            warnings=["implicit_section_without_detected_heading"],
        )
    ]


def page_for_offset(blocks: list[dict], offset: int) -> int | None:
    cursor = 0
    for block in blocks:
        text = str(block.get("text") or "")
        end = cursor + len(text) + 1
        if cursor <= offset <= end:
            value = block.get("page")
            return int(value) if isinstance(value, int) else None
        cursor = end
    return None


def title_decision_for_segment(record: dict[str, Any], parent_text: str, source_file: str) -> tuple[dict[str, Any], dict[str, Any]]:
    if record["segment_type"] == "full_document":
        decision = recover_full_document_title_v2422(record)
        audit = audit_full_document_title_v2422(record, decision)
        return (
            {
                "title": decision.get("title", ""),
                "title_status": audit["status"],
                "title_source": decision.get("source", "missing"),
                "title_recovery_evidence": decision.get("evidence", []),
                "title_recovery_reasons": decision.get("reasons", []),
                "title_hard_issues": audit.get("issues", []),
            },
            audit["consistency"],
        )

    current_title = str(record.get("title") or "")
    decision = recover_title_v242(
        str(record.get("original_title") or ""),
        current_title,
        str(record.get("raw_text") or ""),
        parent_text,
        source_file,
        record["segment_type"],
        record["parent_document_type"],
    )
    title = str(decision.get("title") or "")
    status = str(decision.get("status") or "uncertain")
    source = str(decision.get("source") or "source_record")
    sections = record.get("sections") or []
    consistency = audit_title_consistency_v242(
        document_id=record["document_id"],
        title=title,
        text=str(record.get("raw_text") or ""),
        segment_type=record["segment_type"],
        title_status=status,
        title_source=source,
        sections=sections,
    )
    return (
        {
            "title": title,
            "title_status": status,
            "title_source": source,
            "title_recovery_evidence": decision.get("evidence", []),
            "title_recovery_reasons": [],
            "title_hard_issues": list(decision.get("warnings") or []),
        },
        consistency,
    )


def make_segment(
    *,
    parent: dict[str, Any],
    local_id: str,
    original_title: str,
    local_text: str,
    segment_start: int,
    segment_end: int,
    source_record: dict[str, Any] | None,
    sibling_count: int,
    parent_full: bool,
) -> dict[str, Any]:
    parent_type = parent["parent_document_type"]
    source_file = parent["source_file"]
    base_title = original_title.strip()
    base_type = "full_document" if parent_full else segment_type_v242(
        base_title,
        local_text,
        parent_type,
        sibling_count,
        "uncertain",
        base_title,
        "full_document",
    )
    preliminary = {
        "document_id": local_id,
        "parent_document_id": parent["parent_document_id"],
        "parent_document_type": parent_type,
        "segment_type": base_type,
        "original_title": base_title,
        "title": base_title,
        "title_status": "uncertain",
        "title_source": "source_record" if source_record else "parent_record",
        "raw_text": local_text,
        "segment_text": local_text,
        "extracted_segment_text": local_text,
    }
    final_type, classification_audit = classify_segment_v2421(preliminary, previous_stable_type=base_type)
    preliminary["segment_type"] = final_type
    category_before_title = accident_category_v242(base_title, local_text, final_type, parent.get("accident_category", "other"))
    sections = split_sections_v24(local_text, category_before_title)
    if not sections:
        sections = implicit_section(local_text, base_title, category_before_title)
    preliminary["sections"] = [item.model_dump(mode="json") for item in sections]
    if parent_full:
        preliminary["parent_canonical_title"] = parent.get("canonical_title", "")
        preliminary["document_title"] = parent.get("canonical_title", "")
    if parent_type == "sds":
        try:
            preliminary["sds"] = sds_from_text(local_id, local_text, parent.get("source_url", "")).model_dump(mode="json")
        except Exception as exc:
            preliminary["sds"] = {"review_status": "pending", "extraction_error": f"{type(exc).__name__}:{str(exc)[:200]}"}

    title_fields, consistency = title_decision_for_segment(preliminary, parent.get("normalized_text", ""), source_file)
    preliminary.update(title_fields)
    final_category = accident_category_v242(
        preliminary["title"], local_text, final_type, parent.get("accident_category", "other")
    )
    content_role = content_role_of(parent_type, final_type)
    if final_type == "accident_case_section":
        content_role = "accident_case_reference"
    elif final_type == "guidance_section":
        content_role = parent.get("content_role") or "guidance_reference"
    elif final_type in {"appendix_sds", "embedded_sds"}:
        content_role = "safety_data"
    elif final_type in {"list_fragment", "reference_sentence"}:
        content_role = "context_only"

    section_payload = []
    heading_warnings: list[str] = []
    for section in sections:
        warnings = audit_heading(section.heading, section.content, final_category)
        section_row = section.model_dump(mode="json")
        section_row["warnings"] = list(dict.fromkeys(section_row.get("warnings", []) + warnings))
        section_payload.append(section_row)
        heading_warnings.extend(warnings)

    warnings = list(parent.get("warnings", []))
    warnings.extend(preliminary.get("title_hard_issues", []))
    warnings.extend(heading_warnings)
    if consistency.get("status") == "inconsistent":
        warnings.append("title_body_inconsistent")
    if final_type == "toc_fragment":
        warnings.append("excluded_toc_fragment")
    if source_record is None and not parent_full:
        warnings.append("segment_without_source_record")

    start_page = page_for_offset(parent.get("blocks", []), segment_start)
    end_page = page_for_offset(parent.get("blocks", []), segment_end)
    requires_review = bool(
        parent.get("requires_manual_review")
        or preliminary["title_status"] in {"invalid", "missing", "toc_only"}
        or consistency.get("status") in {"inconsistent", "uncertain"}
        or final_type in {"toc_fragment", "list_fragment", "reference_sentence"}
    )
    quality = "manual_review" if requires_review else "gold_candidate" if parent.get("classification_confidence", 0) >= 0.84 else "silver_candidate"
    result = {
        "document_id": local_id,
        "parent_document_id": parent["parent_document_id"],
        "parent_document_type": parent_type,
        "parent_type_locked": True,
        "parent_type_evidence_scope": parent.get("parent_type_evidence_scope", ""),
        "embedded_content_types": parent.get("embedded_content_types", []),
        "segment_type": final_type,
        "accident_category": final_category,
        "content_role": content_role,
        "title": preliminary["title"],
        "original_title": base_title,
        "title_status": preliminary["title_status"],
        "title_source": preliminary["title_source"],
        "title_recovery_evidence": preliminary.get("title_recovery_evidence", []),
        "title_recovery_reasons": preliminary.get("title_recovery_reasons", []),
        "title_consistency": consistency,
        "classification_reasons": parent.get("classification_reasons", []),
        "classification_evidence": parent.get("classification_evidence", []),
        "classification_confidence": parent.get("classification_confidence", 0),
        "segment_classification_text_source": classification_audit["segment_classification_text_source"],
        "segment_local_text_length": classification_audit["segment_local_text_length"],
        "parent_context_used": classification_audit["parent_context_used"],
        "segment_classification_audit": classification_audit,
        "source_file": source_file,
        "relative_path": parent["relative_path"],
        "source_url": str((source_record or {}).get("url") or parent.get("source_url", "")),
        "source_urls": list(dict.fromkeys(parent.get("source_urls", []) + ([str(source_record.get("url"))] if source_record and source_record.get("url") else []))),
        "source_record_index": (source_record or {}).get("_source_record_index"),
        "source_page_start": start_page,
        "source_page_end": end_page,
        "segment_start": max(0, segment_start),
        "segment_end": max(segment_start, segment_end),
        "raw_text": local_text,
        "normalized_text": normalize_text(local_text),
        "sections": section_payload,
        "warnings": list(dict.fromkeys(warnings)),
        "requires_manual_review": requires_review,
        "quality_grade": quality,
        "review_status": "pending",
        "is_parent_full_document": parent_full,
    }
    if preliminary.get("sds"):
        result["sds"] = preliminary["sds"]
    FullSegmentRecord.model_validate(result)
    return result


def build_segments(parent: dict[str, Any], source_rows: list[dict]) -> list[dict[str, Any]]:
    text = parent.get("normalized_text", "")
    parent_id = parent["parent_document_id"]
    full_id = stable_id("seg", parent_id, "full_document")
    segments = [
        make_segment(
            parent=parent,
            local_id=full_id,
            original_title=parent.get("canonical_title", ""),
            local_text=text,
            segment_start=0,
            segment_end=len(text),
            source_record=None,
            sibling_count=max(1, len(source_rows)),
            parent_full=True,
        )
    ]
    seen: set[tuple[str, str]] = set()
    for order, source in enumerate(source_rows):
        local_text = normalize_text(str(source.get("content") or ""))
        title = str(source.get("title") or "").strip()
        if len(local_text) < 10:
            continue
        fingerprint = (title, hashlib.sha256(local_text.encode("utf-8")).hexdigest())
        if fingerprint in seen:
            continue
        seen.add(fingerprint)
        start = text.find(local_text[: min(len(local_text), 300)]) if text else -1
        if start < 0 and title:
            start = text.find(title)
        if start < 0:
            start = 0
        end = min(len(text), start + len(local_text)) if text else len(local_text)
        local_id = stable_id("rec", parent_id, str(source.get("_source_record_index", order)), title, fingerprint[1])
        segments.append(
            make_segment(
                parent=parent,
                local_id=local_id,
                original_title=title,
                local_text=local_text,
                segment_start=start,
                segment_end=end,
                source_record=source,
                sibling_count=max(1, len(source_rows)),
                parent_full=False,
            )
        )
    return segments


def evidence_item(text: str, document_id: str, source_url: str, section: str, start: int, kind: str, flags: list[str] | None = None) -> dict[str, Any]:
    return {
        "kind": kind,
        "text": text.strip(),
        "document_id": document_id,
        "source_url": source_url,
        "source_section": section,
        "start_position": max(0, start),
        "end_position": max(0, start) + len(text.strip()),
        "flags": list(flags or []),
        "review_status": "pending",
    }


def extract_between(text: str, heading_patterns: list[str], next_heading_pattern: str = r"(?m)^\s*(?:第?\d+|[一二三四五六七八九十]+)[.、．\s]*(?:部分|章|节)?") -> tuple[str, str]:
    for pattern in heading_patterns:
        match = re.search(pattern, text, re.I | re.M)
        if not match:
            continue
        start = match.end()
        next_match = re.search(next_heading_pattern, text[start + 1 :], re.I | re.M)
        end = start + 1 + next_match.start() if next_match else min(len(text), start + 12000)
        return match.group(0).strip(), text[start:end].strip()
    return "", ""


def sentence_evidence(text: str, document_id: str, source_url: str, section: str, kind: str, pattern: str) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    cursor = 0
    for sentence in re.split(r"(?<=[。；;\n])", text):
        stripped = sentence.strip()
        position = text.find(sentence, cursor)
        cursor = max(cursor, position + len(sentence))
        if len(stripped) >= 8 and re.search(pattern, stripped, re.I):
            flags = ["unsafe_or_outdated_candidate"] if re.search(r"催吐|洗胃|小动物试验", stripped) else []
            output.append(evidence_item(stripped, document_id, source_url, section, max(0, position), kind, flags))
    return output[:80]


def enrich_sds(segment: dict[str, Any]) -> dict[str, Any]:
    text = segment.get("normalized_text", "")
    source_url = segment.get("source_url", "")
    record = sds_from_text(segment["document_id"], text, source_url).model_dump(mode="json")
    identity, identity_source = extract_sds_identity({**segment, "sds": record})
    if identity and not record.get("chemical_name_cn"):
        record["chemical_name_cn"] = identity
    cas_values = list(dict.fromkeys(re.findall(r"(?<!\d)(\d{2,7}-\d{2}-\d)(?!\d)", text[:15000])))
    if not record.get("cas") and len(cas_values) == 1:
        record["cas"] = cas_values[0]
    if len(cas_values) > 1:
        record.setdefault("conflict_notes", []).append(f"来源前部识别到多个CAS号：{cas_values[:10]}")
    un_match = re.search(r"(?:UN\s*(?:No\.?|编号|号)?|联合国危险货物编号)\s*[:： ]\s*(\d{4})", text, re.I)
    supplier = re.search(r"(?:供应商|供應商|企业名称|製造商)\s*[:：]\s*([^\n；;]{2,100})", text)
    revision = re.search(r"(?:修订日期|修訂日期|编制日期|製表日期)\s*[:：]?\s*([^\n；;]{4,40})", text)
    record.update({
        "parent_document_id": segment.get("parent_document_id", ""),
        "segment_type": segment.get("segment_type", ""),
        "source_record_index": segment.get("source_record_index"),
        "product_name": identity,
        "identity_source": identity_source,
        "cas_number": record.get("cas", ""),
        "un_number": un_match.group(1) if un_match else "",
        "supplier": supplier.group(1).strip() if supplier else record.get("supplier", ""),
        "revision_date": revision.group(1).strip() if revision else record.get("revision_date", ""),
        "source_file": segment.get("source_file", ""),
        "source_urls": segment.get("source_urls", []),
        "source_sections": [],
        "review_status": "pending",
    })
    section_map = {
        "hazard_classification": ([r"危险性概述", r"危害辨识", r"危害辨識"], r"危险|危害|GHS|分类|類別"),
        "first_aid": ([r"急救措施", r"急救方法"], r"吸入|皮肤|眼睛|食入|就医|醫療"),
        "firefighting_measures": ([r"消防措施", r"灭火措施", r"滅火措施"], r"灭火|消防|燃烧|爆炸"),
        "spill_response": ([r"泄漏应急处理", r"泄漏处理", r"洩漏處理方法"], r"泄漏|疏散|警戒|收容|堵漏"),
        "handling_and_storage": ([r"操作处置与储存", r"安全處置與儲存"], r"操作|储存|貯存|通风"),
        "exposure_controls": ([r"接触控制和个体防护", r"暴露預防措施", r"个体防护"], r"防护|呼吸器|手套|眼镜|通风"),
        "physical_properties": ([r"理化特性", r"物理及化学性质", r"物理及化學性質"], r"外观|沸点|熔点|密度|蒸气"),
        "stability_and_reactivity": ([r"稳定性和反应性", r"安定性及反應性"], r"稳定|禁配|不相容|反应"),
        "toxicology": ([r"毒理学信息", r"毒性資料"], r"毒性|LD50|LC50|刺激"),
        "ecology": ([r"生态学信息", r"生態資料"], r"生态|环境|水生"),
        "disposal": ([r"废弃处置", r"廢棄處置"], r"废弃|处置|容器"),
        "transport": ([r"运输信息", r"運送資料"], r"运输|UN|包装"),
        "regulatory_information": ([r"法规信息", r"法規資料"], r"法规|条例|目录"),
    }
    for field, (headings, filter_pattern) in section_map.items():
        heading, body = extract_between(text, headings)
        if heading and body:
            record["source_sections"].append(heading)
            record[field] = sentence_evidence(body, segment["document_id"], source_url, heading, field, filter_pattern)
    record["requires_manual_review"] = bool(
        not record.get("chemical_name_cn")
        or len(cas_values) > 1
        or record.get("conflict_notes")
        or segment.get("quality_grade") == "manual_review"
    )
    return record


def regulation_record(parent: dict[str, Any]) -> dict[str, Any]:
    text = parent.get("normalized_text", "")
    title = parent.get("title", "") or parent.get("canonical_title", "")
    number = re.search(r"(?:国务院令|部令|主席令|令|文号)\s*(?:第)?\s*([\w〔〕\[\]（）()\-—号]{2,50})", text[:8000])
    issuer = re.search(r"(?m)^\s*((?:中华人民共和国)?(?:国务院|应急管理部|国家安全生产监督管理总局|全国人民代表大会常务委员会|[^\n]{2,20}人民政府))\s*$", text[:6000])
    pub_date = re.search(r"(?:公布日期|发布日期|发布于|印发日期)\s*[:：]?\s*(\d{4}年\d{1,2}月\d{1,2}日)", text[:12000])
    eff_date = re.search(r"(?:施行|生效|实施)(?:日期|时间)?\s*[:：]?\s*(\d{4}年\d{1,2}月\d{1,2}日)|自\s*(\d{4}年\d{1,2}月\d{1,2}日)\s*起施行", text[:16000])
    revision = re.search(r"(?:修订|修正)(?:日期|时间)?\s*[:：]?\s*(\d{4}年\d{1,2}月\d{1,2}日)", text[:16000])
    validity = "unknown"
    if re.search(r"本(?:条例|办法|规定|规程|标准).{0,40}(?:废止|停止执行)|已废止|已经废止", text[:16000]):
        validity = "repealed"
    return {
        "document_id": parent["document_id"],
        "parent_document_id": parent["parent_document_id"],
        "title": title,
        "parent_document_type": "regulation",
        "content_role": "regulatory_reference",
        "document_number": number.group(1).strip() if number else None,
        "issuing_authority": issuer.group(1).strip() if issuer else None,
        "publication_date": pub_date.group(1) if pub_date else None,
        "effective_date": (eff_date.group(1) or eff_date.group(2)) if eff_date else None,
        "expiry_date": None,
        "revision_date": revision.group(1) if revision else None,
        "validity_status": validity,
        "supersedes": None,
        "superseded_by": None,
        "version_notes": "本轮未联网核验；未明确废止/失效时保持unknown。",
        "source_file": parent["source_file"],
        "source_url": parent.get("source_url", ""),
        "source_urls": parent.get("source_urls", []),
        "sha256": parent["sha256"],
        "normalized_text": text,
        "quality_grade": parent.get("quality_grade", "manual_review"),
        "review_status": "pending",
    }


def split_for_chunks(text: str, max_chars: int = 1800, overlap: int = 120) -> list[tuple[int, int, str]]:
    if not text.strip():
        return []
    paragraphs: list[tuple[int, int, str]] = []
    cursor = 0
    for part in re.split(r"(\n\s*\n)", text):
        start = cursor
        cursor += len(part)
        stripped = part.strip()
        if stripped and not re.fullmatch(r"\n\s*\n", part):
            actual = text.find(stripped, start, cursor + 1)
            paragraphs.append((actual if actual >= 0 else start, (actual if actual >= 0 else start) + len(stripped), stripped))
    output: list[tuple[int, int, str]] = []
    buffer: list[str] = []
    chunk_start: int | None = None
    chunk_end = 0
    for start, end, paragraph in paragraphs:
        pieces: list[tuple[int, int, str]] = []
        if len(paragraph) <= max_chars:
            pieces = [(start, end, paragraph)]
        else:
            local_cursor = 0
            for sentence in re.split(r"(?<=[。；;！？!?])", paragraph):
                clean = sentence.strip()
                if not clean:
                    local_cursor += len(sentence)
                    continue
                sentence_start = paragraph.find(clean, local_cursor)
                local_cursor = sentence_start + len(clean)
                if len(clean) <= max_chars:
                    pieces.append((start + sentence_start, start + sentence_start + len(clean), clean))
                else:
                    pos = 0
                    while pos < len(clean):
                        piece_end = min(len(clean), pos + max_chars)
                        pieces.append((start + sentence_start + pos, start + sentence_start + piece_end, clean[pos:piece_end]))
                        pos = max(piece_end - overlap, pos + 1)
        for piece_start, piece_end, piece in pieces:
            projected = len("\n\n".join(buffer + [piece]))
            if buffer and projected > max_chars:
                output.append((chunk_start or 0, chunk_end, "\n\n".join(buffer)))
                buffer = []
                chunk_start = None
            if chunk_start is None:
                chunk_start = piece_start
            buffer.append(piece)
            chunk_end = piece_end
    if buffer:
        output.append((chunk_start or 0, chunk_end, "\n\n".join(buffer)))
    return output


def rag_chunks(segments: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_parent: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for segment in segments:
        by_parent[segment["parent_document_id"]].append(segment)
    selected: list[dict[str, Any]] = []
    for parent_segments in by_parent.values():
        child_candidates = [row for row in parent_segments if row["segment_type"] != "full_document" and row["segment_type"] not in {"toc_fragment", "list_fragment", "reference_sentence"}]
        selected.extend(child_candidates or [row for row in parent_segments if row.get("is_parent_full_document")])
    output: list[dict[str, Any]] = []
    for segment in selected:
        if segment["quality_grade"] == "rejected" or segment["title_consistency"].get("status") == "inconsistent":
            continue
        if not segment.get("source_file"):
            continue
        sections = segment.get("sections") or [{"heading": segment.get("title", ""), "content": segment.get("normalized_text", ""), "start_position": 0}]
        for section_index, section in enumerate(sections):
            content = str(section.get("content") or "").strip()
            base_start = int(section.get("start_position") or 0)
            for piece_index, (start, end, piece) in enumerate(split_for_chunks(content)):
                chunk = {
                    "chunk_id": stable_id("chunk", segment["document_id"], str(section_index), str(piece_index), piece[:80]),
                    "document_id": segment["document_id"],
                    "parent_document_id": segment["parent_document_id"],
                    "title": segment.get("title", ""),
                    "section_title": section.get("heading", ""),
                    "content": piece,
                    "parent_document_type": segment["parent_document_type"],
                    "segment_type": segment["segment_type"],
                    "accident_category": segment["accident_category"],
                    "content_role": segment["content_role"],
                    "chemical_name": (segment.get("sds") or {}).get("chemical_name_cn", ""),
                    "source_file": segment["source_file"],
                    "source_url": segment.get("source_url", ""),
                    "source_page_start": segment.get("source_page_start"),
                    "source_page_end": segment.get("source_page_end"),
                    "char_start": segment["segment_start"] + base_start + start,
                    "char_end": segment["segment_start"] + base_start + end,
                    "quality_grade": segment["quality_grade"],
                    "review_status": "pending",
                    "warnings": segment.get("warnings", []),
                }
                FullRagChunk.model_validate(chunk)
                output.append(chunk)
    return output


def classify_risk(parent: dict[str, Any], segments: list[dict[str, Any]], regulation: dict[str, Any] | None = None, sds: dict[str, Any] | None = None) -> list[str]:
    reasons: list[str] = []
    if parent["parent_document_type"] in {"manual_review", "reject", "other"} or parent.get("requires_manual_review"):
        reasons.append("parent_type_uncertain_or_conflicting")
    if parent.get("parse_status") != "success":
        reasons.append("parse_failure_or_insufficient_text")
    if parent.get("is_scanned_document"):
        reasons.append("scanned_document_requires_ocr_review")
    if not parent.get("source_url"):
        reasons.append("source_url_missing")
    if parent.get("garbled_ratio", 0) > 0.01 or looks_garbled(parent.get("normalized_text", "")):
        reasons.append("garbled_or_short_text")
    for segment in segments:
        if segment["title_status"] in {"missing", "invalid"} and (segment.get("is_parent_full_document") or segment["segment_type"] != "full_document"):
            reasons.append("title_missing_or_invalid")
        if segment["title_consistency"].get("status") == "inconsistent":
            reasons.append("title_body_inconsistent")
        if "multiple_peer_special_headings" in segment["title_consistency"].get("reasons", []):
            reasons.append("multiple_peer_special_headings")
        if HIGH_RISK_CONTENT_RE.search(segment.get("normalized_text", "")[:30000]):
            reasons.append("professional_response_content_requires_review")
    if regulation:
        if regulation.get("validity_status") == "unknown":
            reasons.append("regulation_validity_unknown")
    if sds:
        if not sds.get("chemical_name_cn"):
            reasons.append("sds_identity_unclear")
        if sds.get("conflict_notes"):
            reasons.append("sds_identity_or_cas_conflict")
    return list(dict.fromkeys(reasons))


def validate_rows(manifests: list[dict], parents: list[dict], segments: list[dict], chunks: list[dict]) -> list[dict[str, Any]]:
    checks = [
        ("FullManifestRecord", FullManifestRecord, manifests),
        ("FullParentDocument", FullParentDocument, parents),
        ("FullSegmentRecord", FullSegmentRecord, segments),
        ("FullRagChunk", FullRagChunk, chunks),
    ]
    summaries: list[dict[str, Any]] = []
    for name, model, rows in checks:
        failures: list[dict[str, Any]] = []
        passed = 0
        for index, row in enumerate(rows):
            try:
                model.model_validate(row)
                passed += 1
            except Exception as exc:
                failures.append({"index": index, "document_id": row.get("document_id", row.get("chunk_id", "")), "error": str(exc)[:1000]})
        summaries.append(ValidationSummary(model_name=name, checked=len(rows), passed=passed, failures=failures).model_dump(mode="json"))
    return summaries


def write_review_markdown(path: Path, title: str, rows: list[dict[str, Any]], reason_field: str = "high_risk_reasons") -> None:
    lines = [f"# {title}", "", "所有记录均为自动清洗候选，`review_status=pending`，未进行人工专业批准。", ""]
    for index, row in enumerate(rows, 1):
        text = str(row.get("normalized_text") or row.get("raw_text") or "")
        lines.extend([
            f"## {index}. {row.get('title') or row.get('file_name') or row.get('document_id')}",
            "",
            f"- document_id：`{row.get('document_id', '')}`",
            f"- parent_document_id：`{row.get('parent_document_id', '')}`",
            f"- 父文档类型：`{row.get('parent_document_type', '')}`",
            f"- 子章节类型：`{row.get('segment_type', '')}`",
            f"- 事故类别：`{row.get('accident_category', '')}`",
            f"- 来源文件：`{row.get('source_file', '')}`",
            f"- 来源URL：{row.get('source_url') or '缺失'}",
            f"- 自动警告：{', '.join(row.get(reason_field, []) or row.get('warnings', [])) or '无'}",
            "",
            "关键正文（自动截取）：",
            "",
            "```text",
            text[:1800].replace("```", "` ` `"),
            "```",
            "",
            "人工审核结论：____________________",
            "",
        ])
    path.write_text("\n".join(lines), encoding="utf-8")


def normal_sample_rows(parents: list[dict[str, Any]], high_ids: set[str], rejected_ids: set[str]) -> list[dict[str, Any]]:
    candidates = [row for row in parents if row["document_id"] not in high_ids and row["document_id"] not in rejected_ids]
    if not candidates:
        return []
    rng = random.Random(PILOT_SEED)
    target = max(1, round(len(candidates) * 0.10))
    chosen: dict[str, dict[str, Any]] = {}
    by_type: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in candidates:
        by_type[row["parent_document_type"]].append(row)
    for rows in by_type.values():
        shuffled = list(rows)
        rng.shuffle(shuffled)
        for row in shuffled[: min(3, len(shuffled))]:
            chosen[row["document_id"]] = row
    remaining = [row for row in candidates if row["document_id"] not in chosen]
    rng.shuffle(remaining)
    for row in remaining:
        if len(chosen) >= max(target, len(chosen)):
            break
        chosen[row["document_id"]] = row
    return list(chosen.values())


def report_header(title: str) -> list[str]:
    return [f"# {title}", "", f"生成时间：{utc_now()}", "", "自动处理结果不等于人工专业审核结论；所有记录保持 `pending`。", ""]


def write_reports(
    report_dir: Path,
    *,
    stats: dict[str, Any],
    manifests: list[dict],
    parents: list[dict],
    segments: list[dict],
    chunks: list[dict],
    exact_groups: list[dict],
    near_groups: list[dict],
    regulations: list[dict],
    regulation_conflicts: list[dict],
    sds_rows: list[dict],
    high_rows: list[dict],
    normal_rows: list[dict],
    validation: list[dict],
) -> None:
    parent_counts = Counter(row["parent_document_type"] for row in parents)
    segment_counts = Counter(row["segment_type"] for row in segments)
    parser_counts = Counter(row["parser_selected"] for row in manifests)
    quality_counts = Counter(row.get("quality_grade", "") for row in parents)
    title_counts = Counter(row["title_status"] for row in segments)

    lines = report_header("第一次正式全量清洗运行报告")
    lines.extend([
        "## 运行结论", "",
        f"- 原始文件：{stats['raw_file_count']} 个",
        f"- 成功解析：{stats['parse_success_count']} 个；失败：{stats['parse_failure_count']} 个",
        f"- 独立父文档：{len(parents)} 个；结构化子记录：{len(segments)} 条",
        f"- RAG候选分块：{len(chunks)} 条（未生成Embedding）",
        f"- 高风险审核：{len(high_rows)} 条；普通抽查：{len(normal_rows)} 条",
        f"- raw运行前后签名一致：{stats['raw_unchanged']}",
        f"- Pydantic校验是否全部通过：{stats['pydantic_all_passed']}",
        f"- pytest（运行前/运行后）：{stats.get('pytest_before', '外部执行')}/{stats.get('pytest_after', '待外部执行')}",
        "- 数据库/知识库导入：未执行",
        "- Embedding：未生成",
        "", "## 解析器使用数量", "",
        markdown_table(["解析器", "数量"], sorted(parser_counts.items())),
        "", "## 父文档类型", "",
        markdown_table(["类型", "数量"], sorted(parent_counts.items())),
        "", "## 子章节类型", "",
        markdown_table(["类型", "数量"], sorted(segment_counts.items())),
        "", "## 标题状态", "",
        markdown_table(["状态", "数量"], sorted(title_counts.items())),
        "", "## 质量等级", "",
        markdown_table(["等级", "数量"], sorted(quality_counts.items())),
        "", "## Pydantic校验", "",
        markdown_table(["模型", "检查", "通过", "失败"], ((r["model_name"], r["checked"], r["passed"], len(r["failures"])) for r in validation)),
    ])
    (report_dir / "full_run_report.md").write_text("\n".join(lines), encoding="utf-8")

    classification = report_header("分类统计") + [
        "## 父文档", "", markdown_table(["类型", "数量"], sorted(parent_counts.items())),
        "", "## 子章节", "", markdown_table(["类型", "数量"], sorted(segment_counts.items())),
        "", "## 事故类别", "", markdown_table(["类别", "数量"], sorted(Counter(row["accident_category"] for row in segments).items())),
    ]
    (report_dir / "classification_summary.md").write_text("\n".join(classification), encoding="utf-8")

    failures = [row for row in manifests if row["parse_status"] != "success"]
    parser_report = report_header("解析失败报告") + [
        f"失败文件数量：{len(failures)}", "",
        markdown_table(["文件", "解析器", "错误/警告"], ((r["relative_path"], r["parser_selected"], ", ".join(r.get("parse_error", []) or r.get("warnings", []))) for r in failures)),
    ]
    (report_dir / "parser_failure_report.md").write_text("\n".join(parser_report), encoding="utf-8")

    duplicate_report = report_header("重复资料报告") + [
        f"SHA-256完全重复组：{len(exact_groups)}", "",
        f"近似重复候选关系：{len(near_groups)}", "",
        "重复资料只建立关系，未物理删除原文件。",
    ]
    (report_dir / "duplicate_report.md").write_text("\n".join(duplicate_report), encoding="utf-8")

    title_report = report_header("标题审计报告") + [
        markdown_table(["状态", "数量"], sorted(title_counts.items())), "",
        f"标题正文不一致：{sum(row['title_consistency'].get('status') == 'inconsistent' for row in segments)}",
        f"标题缺失：{title_counts.get('missing', 0)}",
        f"标题无效：{title_counts.get('invalid', 0)}",
    ]
    (report_dir / "title_audit_report.md").write_text("\n".join(title_report), encoding="utf-8")

    boundary_warnings = [(row["document_id"], row["title"], ", ".join(row.get("warnings", []))) for row in segments if any("boundary" in warning or "peer_special" in warning for warning in row.get("warnings", []) + row["title_consistency"].get("reasons", []))]
    boundary_report = report_header("章节边界报告") + [
        f"边界/多同级专项警告：{len(boundary_warnings)}", "",
        markdown_table(["记录", "标题", "警告"], boundary_warnings[:500]),
    ]
    (report_dir / "section_boundary_report.md").write_text("\n".join(boundary_report), encoding="utf-8")

    sds_low = [row for row in sds_rows if row.get("requires_manual_review") or row.get("conflict_notes")]
    sds_quality = report_header("SDS质量报告") + [
        f"SDS记录：{len(sds_rows)}", "",
        f"需人工审核/低质量：{len(sds_low)}", "",
        markdown_table(["记录", "化学品", "CAS", "冲突"], ((r["document_id"], r.get("chemical_name_cn", ""), r.get("cas", ""), "; ".join(r.get("conflict_notes", []))) for r in sds_low)),
    ]
    (report_dir / "sds_quality_report.md").write_text("\n".join(sds_quality), encoding="utf-8")
    (report_dir / "sds_conflict_report.md").write_text("\n".join(report_header("SDS冲突报告") + sds_quality[6:]), encoding="utf-8")

    regulation_candidates = [
        row for row in manifests
        if "/法规/" in f"/{row.get('relative_path', '')}" or row.get("original_type") == "regulation"
    ]
    regulation_candidate_ids = {row["document_id"] for row in regulation_candidates}
    parent_by_id = {row["document_id"]: row for row in parents}
    regulation_rejected = [
        parent_by_id[ident] for ident in regulation_candidate_ids
        if ident in parent_by_id and parent_by_id[ident]["quality_grade"] == "rejected"
    ]
    regulation_preservation = report_header("法律法规保留报告") + [
        f"- 原始法规候选文件数量：{len(regulation_candidates)}",
        f"- 原始法规候选成功解析数量：{sum(r.get('parse_status') == 'success' for r in regulation_candidates)}",
        f"- 原始法规候选解析失败数量：{sum(r.get('parse_status') != 'success' for r in regulation_candidates)}",
        f"- 分类为regulation的数量：{parent_counts.get('regulation', 0)}",
        f"- 分类为guidance_reference的数量：{parent_counts.get('guidance_reference', 0)}",
        f"- 完全重复组数量：{sum(any(i in {r['document_id'] for r in regulations} for i in g.get('duplicate_record_ids', [])) for g in exact_groups)}",
        f"- 近似重复候选数量：{sum(any(i in {r['document_id'] for r in regulations} for i in g.get('document_ids', [])) for g in near_groups)}",
        f"- 新旧版本冲突数量：{len(regulation_conflicts)}",
        f"- validity_status=unknown数量：{sum(r.get('validity_status') == 'unknown' for r in regulations)}",
        f"- 进入人工审核数量：{sum(r.get('validity_status') in {'unknown', 'conflict_requires_review'} for r in regulations)}",
        f"- 进入拒绝队列数量：{len(regulation_rejected)}",
        "- 是否存在因“发布时间旧”而被删除的法规：否",
        "", "拒绝法规仅保留状态和原因；未删除任何原文件。",
    ]
    (report_dir / "regulation_preservation_report.md").write_text("\n".join(regulation_preservation), encoding="utf-8")
    version_report = report_header("法规版本报告") + [
        f"版本冲突候选组：{len(regulation_conflicts)}", "",
        markdown_table(["组", "规范标题", "记录数", "说明"], ((r["conflict_id"], r["canonical_title"], len(r["document_ids"]), r["reason"]) for r in regulation_conflicts)),
    ]
    (report_dir / "regulation_version_report.md").write_text("\n".join(version_report), encoding="utf-8")

    traceable = sum(bool(row.get("source_file")) and Path(row["source_file"]).exists() for row in parents)
    trace_report = report_header("来源追溯报告") + [
        f"父文档来源文件可追溯：{traceable}/{len(parents)}（{round(traceable/max(1,len(parents))*100,2)}%）",
        f"具备来源URL：{sum(bool(row.get('source_url')) for row in parents)}/{len(parents)}",
        f"来源URL缺失记录仍保留并进入审核：{sum(not bool(row.get('source_url')) for row in parents)}",
    ]
    (report_dir / "source_traceability_report.md").write_text("\n".join(trace_report), encoding="utf-8")
    quality_report = report_header("质量分布报告") + [markdown_table(["等级", "数量"], sorted(quality_counts.items()))]
    (report_dir / "quality_distribution_report.md").write_text("\n".join(quality_report), encoding="utf-8")

    write_review_markdown(report_dir / "high_risk_review.md", "高风险审核报告", high_rows)
    enterprise_rows = [row for row in parents if row["parent_document_type"] in {"enterprise_plan", "enterprise_group_plan"}]
    write_review_markdown(report_dir / "enterprise_plan_candidates_review.md", "企业预案候选审核包", enterprise_rows)
    write_review_markdown(report_dir / "sds_conflict_and_low_quality_review.md", "SDS冲突与低质量审核包", sds_low)
    write_review_markdown(report_dir / "normal_records_sample_review.md", "普通结果10%抽查包", normal_rows)

    excluded_types = Counter(row["segment_type"] for row in segments if row["segment_type"] in {"toc_fragment", "list_fragment", "reference_sentence"})
    rag_report = report_header("RAG分块准备报告") + [
        f"- 候选分块：{len(chunks)}",
        "- Embedding：未生成",
        f"- 来源文件可追溯率：{round(sum(Path(c['source_file']).exists() for c in chunks)/max(1,len(chunks))*100,2)}%",
        f"- 待审核状态：{sum(c['review_status'] == 'pending' for c in chunks)}/{len(chunks)}",
        "", "默认排除但保留的片段：", "",
        markdown_table(["类型", "数量"], sorted(excluded_types.items())),
    ]
    (report_dir / "rag_chunk_readiness_report.md").write_text("\n".join(rag_report), encoding="utf-8")


def regulation_version_conflicts(regulations: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_title: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in regulations:
        title = re.sub(r"[（(](?:20\d{2}年)?(?:修订|修正|版)[）)]", "", row.get("title", ""))
        key = re.sub(r"\s+", "", title).strip("《》")
        if key:
            by_title[key].append(row)
    conflicts: list[dict[str, Any]] = []
    for key, rows in by_title.items():
        versions = {(row.get("document_number"), row.get("publication_date"), row.get("revision_date")) for row in rows}
        if len(rows) > 1 and len(versions) > 1:
            conflicts.append({
                "conflict_id": stable_id("reg-conflict", key),
                "canonical_title": key,
                "document_ids": [row["document_id"] for row in rows],
                "versions": [
                    {"document_id": row["document_id"], "document_number": row.get("document_number"), "publication_date": row.get("publication_date"), "revision_date": row.get("revision_date"), "validity_status": "conflict_requires_review"}
                    for row in rows
                ],
                "reason": "同名法规存在不同文号/发布日期/修订日期，未联网核验有效性",
                "review_status": "pending",
            })
    return conflicts


def write_manifest_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fields = [
        "document_id", "source_file", "relative_path", "file_name", "file_extension", "file_size",
        "sha256", "modified_time", "parser_selected", "source_url", "source_url_status", "parse_status",
        "parse_error", "page_count", "text_length", "is_scanned_document", "review_status",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            output = dict(row)
            output["parse_error"] = "; ".join(row.get("parse_error", []) or row.get("warnings", []))
            writer.writerow(output)


def update_final_pytest(root: Path, value: str) -> int:
    manifest_path = root / OUTPUT_NAMES["manifests"] / "run_manifest.json"
    report_path = root / OUTPUT_NAMES["reports"] / "full_run_report.md"
    if not manifest_path.exists() or not report_path.exists():
        raise FileNotFoundError("full run outputs do not exist")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["pytest_after"] = value
    manifest["pytest_after_passed"] = "passed" in value and "failed" not in value
    if isinstance(manifest.get("stats"), dict):
        manifest["stats"]["pytest_after"] = value
    json_dump(manifest_path, manifest)
    checkpoint_path = root / OUTPUT_NAMES["checkpoints"] / "full_run_complete.json"
    if checkpoint_path.exists():
        checkpoint = json.loads(checkpoint_path.read_text(encoding="utf-8"))
        if isinstance(checkpoint.get("stats"), dict):
            checkpoint["stats"]["pytest_after"] = value
        json_dump(checkpoint_path, checkpoint)
    report = report_path.read_text(encoding="utf-8")
    report = re.sub(r"- pytest（运行前/运行后）：[^\n]+", f"- pytest（运行前/运行后）：{manifest.get('pytest_before', '已通过')}/{value}", report)
    report_path.write_text(report, encoding="utf-8")
    return 0


def run_full(root: Path, raw_dir: Path, resume: bool = False) -> int:
    if not raw_dir.is_dir():
        print(json.dumps({"status": "stopped", "reason": "raw_directory_missing", "checked_path": str(raw_dir)}, ensure_ascii=False), flush=True)
        return 2

    # Baseline is completed in memory before any full-run output directory is created.
    baseline_rows = scan_raw(raw_dir)
    baseline_sig = baseline_signature(baseline_rows)
    paths = {name: root / dirname for name, dirname in OUTPUT_NAMES.items()}
    existing = [str(path) for path in paths.values() if path.exists()]
    if existing and not resume:
        print(json.dumps({"status": "stopped", "reason": "output_directory_already_exists", "paths": existing}, ensure_ascii=False), flush=True)
        return 3
    for path in paths.values():
        path.mkdir(parents=True, exist_ok=True)
    parsed_dir = paths["checkpoints"] / "parsed_documents"
    parsed_dir.mkdir(parents=True, exist_ok=True)
    log_path = paths["logs"] / "full_run.log.jsonl"

    def log(stage: str, message: str, **extra: Any) -> None:
        payload = {"time": utc_now(), "stage": stage, "message": message, **extra}
        with log_path.open("a", encoding="utf-8") as stream:
            stream.write(json.dumps(payload, ensure_ascii=False) + "\n")

    source_map, unmatched_sources = load_source_records(root, baseline_rows)
    meta_by_id: dict[str, dict[str, Any]] = {}
    for row in baseline_rows:
        metadata = source_metadata(source_map.get(row["document_id"], []))
        if not metadata["source_urls"]:
            provenance_urls = local_where_from_urls(Path(row["source_file"]))
            if provenance_urls:
                metadata["source_urls"] = provenance_urls
                metadata["source_url"] = provenance_urls[0]
                metadata["source_url_status"] = "present_from_local_file_provenance"
        meta_by_id[row["document_id"]] = metadata
    for row in baseline_rows:
        row.update(meta_by_id[row["document_id"]])
        row.update({
            "parser_selected": "pending",
            "parse_status": "pending",
            "parse_error": [],
            "page_count": 0,
            "text_length": 0,
            "is_scanned_document": False,
            "review_status": "pending",
        })
    exact_groups = exact_duplicate_groups(baseline_rows)
    duplicate_of: dict[str, str] = {}
    for group in exact_groups:
        for ident in group["duplicate_record_ids"]:
            duplicate_of[ident] = group["canonical_record_id"]
    for row in baseline_rows:
        row["duplicate_of"] = duplicate_of.get(row["document_id"])

    initial_run_manifest = {
        "run_id": stable_id("full-run", baseline_sig, utc_now()),
        "version": "full_v1_using_v2.4.2.2_preflight",
        "started_at": utc_now(),
        "root_dir": str(root),
        "raw_dir": str(raw_dir),
        "raw_file_count_before": len(baseline_rows),
        "raw_baseline_signature": baseline_sig,
        "fixed_seed": PILOT_SEED,
        "review_status_policy": "pending_only",
        "database_import_executed": False,
        "knowledge_base_import_executed": False,
        "embedding_generated": False,
        "raw_mutation_authorized": False,
        "pytest_before": os.environ.get("FULL_RUN_PYTEST_BEFORE", "passed_before_runner"),
        "tools": {
            name: package_version(package)
            for name, package in {
                "docling": "docling", "datasketch": "datasketch", "rapidfuzz": "RapidFuzz",
                "pydantic": "pydantic", "presidio-analyzer": "presidio-analyzer", "pytest": "pytest",
                "pypdf": "pypdf", "pdfplumber": "pdfplumber", "python-docx": "python-docx",
            }.items()
        },
    }
    json_dump(paths["manifests"] / "run_manifest.json", initial_run_manifest)
    json_dump(paths["manifests"] / "sha256_manifest.json", {row["relative_path"]: row["sha256"] for row in baseline_rows})
    json_dump(paths["manifests"] / "source_url_manifest.json", [
        {"document_id": row["document_id"], "source_file": row["source_file"], "source_urls": row["source_urls"], "source_url_status": row["source_url_status"]}
        for row in baseline_rows
    ])
    json_dump(paths["manifests"] / "unmatched_source_records.json", unmatched_sources)
    jsonl_dump(paths["manifests"] / "raw_file_manifest.jsonl", baseline_rows)
    write_manifest_csv(paths["manifests"] / "raw_file_manifest.csv", baseline_rows)
    log("baseline", "raw baseline captured", count=len(baseline_rows), signature=baseline_sig)

    parsed_by_id: dict[str, dict[str, Any]] = {}
    # A small PDF goes first so the stable parser exercises Docling once before audited fallbacks.
    parse_order = sorted(baseline_rows, key=lambda row: (row["file_extension"] != "pdf", row["file_size"], row["relative_path"]))
    print(f"FULL_RUN_STAGE parse_start files={len(parse_order)}", flush=True)
    for index, row in enumerate(parse_order, 1):
        checkpoint = parsed_dir / f"{row['document_id']}.json"
        retry_html_metadata_bug = False
        if resume and checkpoint.exists():
            payload = json.loads(checkpoint.read_text(encoding="utf-8"))
            retry_html_metadata_bug = bool(
                row["file_extension"] in {"html", "htm"}
                and payload.get("parse_status") == "failed"
                and any("PackageNotFoundError" in warning and "html_local" in warning for warning in payload.get("warnings", []))
            )
        if not (resume and checkpoint.exists()) or retry_html_metadata_bug:
            with parse_deadline(120):
                payload = parse_file(row, source_map.get(row["document_id"], []))
            json_dump(checkpoint, payload)
        parsed_by_id[row["document_id"]] = payload
        row.update({
            "parser_selected": payload["parser"],
            "parse_status": payload["parse_status"],
            "parse_error": payload.get("warnings", []),
            "warnings": payload.get("warnings", []),
            "page_count": payload.get("page_count", 0),
            "text_length": payload.get("text_length", 0),
            "is_scanned_document": payload.get("is_scanned_document", False),
        })
        if index % 10 == 0 or index == len(parse_order):
            print(f"FULL_RUN_PROGRESS parsed={index}/{len(parse_order)} success={sum(item['parse_status']=='success' for item in baseline_rows)}", flush=True)
            log("parse", "batch progress", processed=index, total=len(parse_order))
    jsonl_dump(paths["manifests"] / "raw_file_manifest.jsonl", baseline_rows)
    write_manifest_csv(paths["manifests"] / "raw_file_manifest.csv", baseline_rows)
    json_dump(paths["checkpoints"] / "parse_stage_complete.json", {"completed_at": utc_now(), "count": len(parsed_by_id)})

    print("FULL_RUN_STAGE classification_start", flush=True)
    parents: list[dict[str, Any]] = []
    parent_by_id: dict[str, dict[str, Any]] = {}
    for index, manifest in enumerate(baseline_rows, 1):
        payload = parsed_by_id[manifest["document_id"]]
        text = payload.get("normalized_text", "")
        source_rows = source_map.get(manifest["document_id"], [])
        source_title = choose_parent_title(Path(manifest["source_file"]), text, source_rows)
        lock = lock_parent_type_v241(manifest["document_id"], source_title, text, manifest.get("original_type", ""))
        classified = lock.result.model_dump(mode="json")
        parent_type = classified["parent_document_type"]
        if parent_type not in PARENT_TYPES:
            parent_type = "manual_review"
            classified["requires_manual_review"] = True
        quality = "rejected" if manifest["parse_status"] != "success" else "manual_review" if parent_type in {"manual_review", "reject", "other"} or classified.get("requires_manual_review") else "gold_candidate" if classified.get("classification_confidence", 0) >= 0.84 else "silver_candidate"
        canonical_title = lock.cover_title or source_title
        parent = {
            "record_level": "parent_document",
            "document_id": manifest["document_id"],
            "parent_document_id": manifest["document_id"],
            "parent_document_type": parent_type,
            "parent_type_locked": True,
            "parent_type_evidence_scope": lock.evidence_scope,
            "embedded_content_types": lock.embedded_content_types,
            "canonical_title": canonical_title,
            "title": canonical_title,
            "classification_reasons": classified.get("classification_reasons", []),
            "classification_evidence": classified.get("classification_evidence", []),
            "classification_confidence": classified.get("classification_confidence", 0),
            "requires_manual_review": classified.get("requires_manual_review", True),
            "accident_category": classified.get("accident_category", "other"),
            "content_role": classified.get("content_role", content_role_of(parent_type, "full_document")),
            "source_file": manifest["source_file"],
            "relative_path": manifest["relative_path"],
            "source_url": manifest["source_url"],
            "source_urls": manifest["source_urls"],
            "source_url_status": manifest["source_url_status"],
            "sha256": manifest["sha256"],
            "duplicate_of": manifest.get("duplicate_of"),
            "parser_selected": payload["parser"],
            "parser_version": payload.get("parser_version", ""),
            "parse_status": payload["parse_status"],
            "parse_warnings": payload.get("warnings", []),
            "page_count": payload.get("page_count", 0),
            "text_length": payload.get("text_length", 0),
            "chinese_ratio": payload.get("chinese_ratio", 0),
            "garbled_ratio": payload.get("garbled_ratio", 0),
            "is_scanned_document": payload.get("is_scanned_document", False),
            "blocks": payload.get("blocks", []),
            "raw_text": payload.get("raw_text", ""),
            "normalized_text": text,
            "warnings": list(payload.get("warnings", [])),
            "quality_grade": quality,
            "review_status": "pending",
        }
        if parent_type in {"enterprise_plan", "enterprise_group_plan", "organization_plan", "government_or_regional_plan"} and text:
            try:
                redacted, audits = redact(parent["document_id"], text[:50000])
                parent["redacted_candidate_preview"] = redacted[:12000]
                parent["redaction_audit_count"] = len(audits)
                parent["redaction_audit"] = [item.model_dump(mode="json") for item in audits]
            except Exception as exc:
                parent["warnings"].append(f"redaction_failed:{type(exc).__name__}")
                parent["redaction_audit"] = []
        else:
            parent["redaction_audit"] = []
        FullParentDocument.model_validate(parent)
        parents.append(parent)
        parent_by_id[parent["document_id"]] = parent
        if index % 50 == 0:
            print(f"FULL_RUN_PROGRESS classified={index}/{len(baseline_rows)}", flush=True)

    near_groups = near_duplicate_pairs(parents)
    json_dump(paths["checkpoints"] / "classification_stage_complete.json", {"completed_at": utc_now(), "parents": len(parents), "near_pairs": len(near_groups)})

    print("FULL_RUN_STAGE segment_and_title_start", flush=True)
    segments: list[dict[str, Any]] = []
    segment_errors: list[dict[str, Any]] = []
    for index, parent in enumerate(parents, 1):
        try:
            built = build_segments(parent, source_map.get(parent["document_id"], []))
            full_segment = next((item for item in built if item["segment_type"] == "full_document"), None)
            if full_segment:
                parent["title"] = full_segment["title"]
                parent["title_status"] = full_segment["title_status"]
                parent["title_source"] = full_segment["title_source"]
                parent["title_consistency"] = full_segment["title_consistency"]
            for segment in built:
                try:
                    model_sections = [PlanSection.model_validate(item) for item in segment.get("sections", [])]
                    segment["knowledge_items"] = [item.model_dump(mode="json") for item in extract_knowledge(segment["document_id"], segment.get("source_url", ""), model_sections)]
                except Exception as exc:
                    segment["warnings"].append(f"knowledge_extraction_failed:{type(exc).__name__}")
                    segment["knowledge_items"] = []
            segments.extend(built)
        except Exception as exc:
            segment_errors.append({"parent_document_id": parent["document_id"], "error": f"{type(exc).__name__}:{str(exc)[:1000]}", "review_status": "pending"})
            parent["warnings"].append("segment_generation_failed")
            parent["requires_manual_review"] = True
            parent["quality_grade"] = "manual_review"
        if index % 25 == 0 or index == len(parents):
            print(f"FULL_RUN_PROGRESS segmented_parents={index}/{len(parents)} segments={len(segments)} errors={len(segment_errors)}", flush=True)
    json_dump(paths["checkpoints"] / "segment_stage_complete.json", {"completed_at": utc_now(), "segments": len(segments), "errors": segment_errors})

    print("FULL_RUN_STAGE sds_regulation_rag_start", flush=True)
    segments_by_parent: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for segment in segments:
        segments_by_parent[segment["parent_document_id"]].append(segment)
    sds_rows: list[dict[str, Any]] = []
    for segment in segments:
        if (segment["parent_document_type"] == "sds" and segment["segment_type"] == "full_document") or segment["segment_type"] in {"appendix_sds", "embedded_sds"}:
            try:
                enriched = enrich_sds(segment)
                segment["sds"] = enriched
                sds_rows.append(enriched)
            except Exception as exc:
                segment["warnings"].append(f"sds_extraction_failed:{type(exc).__name__}")
                segment["requires_manual_review"] = True
    regulations = [regulation_record(parent) for parent in parents if parent["parent_document_type"] == "regulation"]
    regulation_conflicts = regulation_version_conflicts(regulations)
    chunks = rag_chunks(segments)

    high_rows: list[dict[str, Any]] = []
    manual_rows: list[dict[str, Any]] = []
    rejected_rows: list[dict[str, Any]] = []
    sds_by_parent = {row.get("parent_document_id", row["document_id"]): row for row in sds_rows}
    regulation_by_parent = {row["document_id"]: row for row in regulations}
    for parent in parents:
        reasons = classify_risk(parent, segments_by_parent.get(parent["document_id"], []), regulation_by_parent.get(parent["document_id"]), sds_by_parent.get(parent["document_id"]))
        if reasons:
            review = dict(parent)
            review["high_risk_reasons"] = reasons
            high_rows.append(review)
        if parent.get("requires_manual_review") or reasons:
            manual_rows.append({
                "document_id": parent["document_id"],
                "parent_document_id": parent["parent_document_id"],
                "title": parent.get("title", ""),
                "source_file": parent["source_file"],
                "source_url": parent.get("source_url", ""),
                "stage": "full_run_v1",
                "reasons": reasons or parent.get("classification_reasons", []),
                "review_status": "pending",
            })
        if parent["quality_grade"] == "rejected":
            rejected_rows.append({
                "document_id": parent["document_id"],
                "parent_document_id": parent["parent_document_id"],
                "source_file": parent["source_file"],
                "source_url": parent.get("source_url", ""),
                "sha256": parent["sha256"],
                "reject_reason": parent.get("parse_warnings", []) or parent.get("classification_reasons", []),
                "raw_file_retained": True,
                "review_status": "pending",
            })
    rejected_ids = {row["document_id"] for row in rejected_rows}
    high_ids = {row["document_id"] for row in high_rows}
    normal_candidates = [
        row for row in segments
        if parent_by_id[row["parent_document_id"]]["quality_grade"] != "rejected"
        and row["parent_document_type"] not in {"manual_review", "reject", "other"}
        and row["segment_type"] not in {"toc_fragment", "list_fragment", "reference_sentence"}
        and (row["segment_type"] != "full_document" or row.get("is_parent_full_document"))
        and row["title_status"] in {"valid", "uncertain"}
        and row["title_consistency"].get("status") != "inconsistent"
        and not HIGH_RISK_CONTENT_RE.search(row.get("normalized_text", "")[:30000])
    ]
    for row in normal_candidates:
        row["normal_sample_eligibility"] = "non_high_risk_non_rejected_structured_record"
    normal_rows = normal_sample_rows(normal_candidates, set(), set())

    validation = validate_rows(baseline_rows, parents, segments, chunks)
    pydantic_all_passed = all(not row["failures"] for row in validation)
    json_dump(paths["checkpoints"] / "pydantic_validation.json", validation)

    cleaned = paths["cleaned"]
    all_records = [{**row, "record_level": "parent_document"} for row in parents] + [{**row, "record_level": "segment"} for row in segments]
    output_map = {
        "all_records.jsonl": all_records,
        "parent_documents.jsonl": parents,
        "segment_records.jsonl": segments,
        "rag_chunks.jsonl": chunks,
        "enterprise_plans.jsonl": [row for row in parents if row["parent_document_type"] == "enterprise_plan"],
        "enterprise_group_plans.jsonl": [row for row in parents if row["parent_document_type"] == "enterprise_group_plan"],
        "government_plans.jsonl": [row for row in parents if row["parent_document_type"] == "government_or_regional_plan"],
        "organization_plans.jsonl": [row for row in parents if row["parent_document_type"] == "organization_plan"],
        "regulations.jsonl": regulations,
        "guidance_references.jsonl": [row for row in parents if row["parent_document_type"] == "guidance_reference"],
        "accident_cases.jsonl": [row for row in parents if row["parent_document_type"] == "accident_case"],
        "chemical_catalogs.jsonl": [row for row in parents if row["parent_document_type"] == "chemical_catalog"],
        "sds_records.jsonl": [row for row in sds_rows if parent_by_id.get(row.get("parent_document_id", ""), {}).get("parent_document_type") == "sds" and row.get("segment_type") == "full_document" and row.get("source_record_index") is None],
        "appendix_sds_records.jsonl": [row for row in sds_rows if row.get("segment_type") in {"appendix_sds", "embedded_sds"}],
        "onsite_disposal_plans.jsonl": [row for row in segments if row["segment_type"] == "onsite_disposal_plan"],
        "embedded_special_sections.jsonl": [row for row in segments if row["segment_type"] == "embedded_special_section"],
        "regulation_manual_review.jsonl": [row for row in regulations if row["validity_status"] in {"unknown", "conflict_requires_review"}],
        "manual_review_queue.jsonl": manual_rows,
        "high_risk_review_queue.jsonl": high_rows,
        "rejected_records.jsonl": rejected_rows,
        "excluded_but_retained_records.jsonl": [row for row in segments if row["segment_type"] in {"toc_fragment", "list_fragment", "reference_sentence"} or row["quality_grade"] == "rejected" or row["title_consistency"].get("status") == "inconsistent"],
    }
    for name, rows in output_map.items():
        jsonl_dump(cleaned / name, rows)
    json_dump(cleaned / "duplicate_groups.json", exact_groups)
    json_dump(cleaned / "near_duplicate_groups.json", near_groups)
    regulation_ids = {row["document_id"] for row in regulations}
    json_dump(cleaned / "regulation_duplicates.json", [group for group in exact_groups if group["canonical_record_id"] in regulation_ids or any(item in regulation_ids for item in group["duplicate_record_ids"])])
    json_dump(cleaned / "regulation_version_conflicts.json", regulation_conflicts)
    jsonl_dump(cleaned / "redaction_audit.jsonl", (audit for parent in parents for audit in parent.get("redaction_audit", [])))
    jsonl_dump(cleaned / "knowledge_items.jsonl", (item for segment in segments for item in segment.get("knowledge_items", [])))

    final_raw_rows = scan_raw(raw_dir)
    final_sig = baseline_signature(final_raw_rows)
    raw_unchanged = baseline_sig == final_sig and len(final_raw_rows) == len(baseline_rows)
    source_traceable = sum(Path(row["source_file"]).exists() for row in parents) / max(1, len(parents))
    stats = {
        "raw_file_count": len(baseline_rows),
        "extension_counts": dict(sorted(Counter(row["file_extension"] or "[no_extension]" for row in baseline_rows).items())),
        "parse_success_count": sum(row["parse_status"] == "success" for row in baseline_rows),
        "parse_failure_count": sum(row["parse_status"] != "success" for row in baseline_rows),
        "exact_duplicate_group_count": len(exact_groups),
        "exact_duplicate_file_count": sum(len(group["duplicate_record_ids"]) for group in exact_groups),
        "near_duplicate_relation_count": len(near_groups),
        "parent_document_count": len(parents),
        "structured_record_count": len(segments),
        "rag_chunk_count": len(chunks),
        "parent_type_counts": dict(sorted(Counter(row["parent_document_type"] for row in parents).items())),
        "segment_type_counts": dict(sorted(Counter(row["segment_type"] for row in segments).items())),
        "sds_count": sum(row["parent_document_type"] == "sds" for row in parents),
        "sds_attachment_count": sum(row["segment_type"] in {"appendix_sds", "embedded_sds"} for row in segments),
        "high_risk_count": len(high_rows),
        "normal_sample_count": len(normal_rows),
        "rejected_count": len(rejected_rows),
        "regulation_version_conflict_count": len(regulation_conflicts),
        "regulation_validity_unknown_count": sum(row["validity_status"] == "unknown" for row in regulations),
        "source_traceability_rate": round(source_traceable * 100, 4),
        "raw_sha256_available_count": sum(bool(row.get("sha256")) for row in baseline_rows),
        "raw_sha256_unavailable_count": sum(not bool(row.get("sha256")) for row in baseline_rows),
        "raw_sha256_complete": all(bool(row.get("sha256")) for row in baseline_rows),
        "all_pending": all(row.get("review_status") == "pending" for row in all_records + chunks + sds_rows + regulations),
        "raw_unchanged": raw_unchanged,
        "pydantic_all_passed": pydantic_all_passed,
        "pytest_before": initial_run_manifest["pytest_before"],
        "pytest_after": "pending_external_final_test",
        "database_import_executed": False,
        "embedding_generated": False,
    }
    write_reports(
        paths["reports"],
        stats=stats,
        manifests=baseline_rows,
        parents=parents,
        segments=segments,
        chunks=chunks,
        exact_groups=exact_groups,
        near_groups=near_groups,
        regulations=regulations,
        regulation_conflicts=regulation_conflicts,
        sds_rows=sds_rows,
        high_rows=high_rows,
        normal_rows=normal_rows,
        validation=validation,
    )
    run_manifest = {
        **initial_run_manifest,
        "completed_at": utc_now(),
        "raw_file_count_after": len(final_raw_rows),
        "raw_final_signature": final_sig,
        "raw_unchanged": raw_unchanged,
        "stats": stats,
        "output_files": sorted(str(path) for directory in paths.values() for path in directory.rglob("*") if path.is_file()),
        "pydantic_validation": validation,
        "pytest_after": "pending_external_final_test",
    }
    json_dump(paths["manifests"] / "run_manifest.json", run_manifest)
    json_dump(paths["checkpoints"] / "full_run_complete.json", {"completed_at": utc_now(), "stats": stats})
    print("FULL_RUN_COMPLETE " + json.dumps(stats, ensure_ascii=False), flush=True)
    if not raw_unchanged or not stats["raw_sha256_complete"] or not pydantic_all_passed or not stats["all_pending"]:
        return 4
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--raw", type=Path, required=True)
    parser.add_argument("--resume", action="store_true")
    parser.add_argument("--record-final-pytest", default="")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    root = args.root.expanduser().resolve()
    raw = args.raw.expanduser().resolve()
    if args.record_final_pytest:
        return update_final_pytest(root, args.record_final_pytest)
    return run_full(root, raw, args.resume)


if __name__ == "__main__":
    raise SystemExit(main())
